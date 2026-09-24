"""Exports a collection of pins as Markdown or a Word document, with citations."""
import io
from datetime import date
from typing import List

from docx import Document as WordDocument
from docx.shared import Pt

from src.models import Pin


def _md_cell(text: str) -> str:
    return text.replace("|", "\\|").replace("\n", " ")


def to_markdown(collection_name: str, pins: List[Pin]) -> str:
    lines = [f"# {collection_name}", "", f"_Exported {date.today().isoformat()} · {len(pins)} items_", ""]
    for i, pin in enumerate(pins, 1):
        lines += [f"## {i}. {pin.citation}", ""]
        if pin.table:
            if pin.table.get("caption"):
                lines += [f"*{pin.table['caption']}*", ""]
            columns = pin.table["columns"]
            lines.append("| " + " | ".join(_md_cell(c) for c in columns) + " |")
            lines.append("|" + "---|" * len(columns))
            lines += ["| " + " | ".join(_md_cell(c) for c in row) + " |" for row in pin.table["rows"]]
        else:
            lines.append(f"> {pin.text}")
        lines.append("")
        if pin.note:
            lines += [f"**Note:** {pin.note}", ""]
        if pin.query:
            lines += [f"_Found by searching:_ `{pin.query}`", ""]
    return "\n".join(lines)


def to_docx(collection_name: str, pins: List[Pin]) -> bytes:
    doc = WordDocument()
    doc.add_heading(collection_name, level=0)
    doc.add_paragraph(f"Exported {date.today().isoformat()} · {len(pins)} items")

    for i, pin in enumerate(pins, 1):
        doc.add_heading(f"{i}. {pin.citation}", level=2)
        if pin.table:
            if pin.table.get("caption"):
                doc.add_paragraph().add_run(pin.table["caption"]).italic = True
            columns = pin.table["columns"]
            table = doc.add_table(rows=1, cols=len(columns), style="Table Grid")
            for cell, name in zip(table.rows[0].cells, columns):
                cell.text = name
                for run in cell.paragraphs[0].runs:
                    run.bold = True
            for row in pin.table["rows"]:
                cells = table.add_row().cells
                for cell, value in zip(cells, row):
                    cell.text = value
        else:
            doc.add_paragraph(pin.text, style="Quote")
        if pin.note:
            note = doc.add_paragraph()
            note.add_run("Note: ").bold = True
            note.add_run(pin.note)
        if pin.query:
            found = doc.add_paragraph()
            found_run = found.add_run(f"Found by searching: {pin.query}")
            found_run.italic = True
            found_run.font.size = Pt(9)

    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()
